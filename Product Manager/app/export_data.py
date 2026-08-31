"""
Exports one or more sourcer xlsx outputs + readme_summaries.json to app/data.json.js
for the local viewer (index.html). Run manually after each sourcer run:
  ../.venv/bin/python fetch_readmes.py   (optional but recommended, fills readme_summary)
  ../.venv/bin/python export_data.py

Uses openpyxl directly (not pandas) for repo_url/owner_url/owner_blog/owner_x/owner_linkedin
because write_excel_with_fallback() in the sourcer scripts replaces those cells' visible
text with a display label ("Repo", "LinkedIn", ...) and stores the real URL only as a
hyperlink target — pandas.read_excel only sees the label, not the link.

Sources missing on disk (e.g. a sourcer variant that hasn't been run yet) are skipped.
"""
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).parent.parent
README_CACHE = Path(__file__).parent / "readme_summaries.json"
OUT = Path(__file__).parent / "data.json.js"

# One entry per sourcer script's output. "source"/"source_label" tag each row so the
# viewer can filter by which sourcer found it.
SOURCES = [
    {
        "source": "react",
        "source_label": "FDE / React (fde_react_sourcer.py)",
        "xlsx": ROOT / "github_fde_react_candidates.xlsx",
    },
    {
        "source": "grok",
        "source_label": "FDE / Product-Design Engineer (fde_grok_react_sourcer.py)",
        "xlsx": ROOT / "github_fde_grok_candidates.xlsx",
    },
]

# Columns whose cell value was replaced by a display label; read the hyperlink target instead.
HYPERLINK_COLUMNS = {"repo_url", "owner_url", "owner_blog", "owner_x", "owner_linkedin"}

COLUMNS = [
    "owner_login", "owner_name", "owner_url", "owner_bio", "owner_location",
    "owner_country", "owner_city", "owner_email", "owner_blog", "owner_x",
    "owner_linkedin", "owner_extra_links",
    "repo_full_name", "repo_url", "description", "language", "stars", "forks",
    "created_at", "pushed_at", "matched_query", "scan_kind",
    "contributors_top", "contributors_top_n",
]


def clean(v):
    if v is None:
        return ""
    return v


def load_rows(xlsx_path: Path, source: str, source_label: str) -> list[dict]:
    wb = load_workbook(xlsx_path)
    ws = wb.active
    header = {c.value: i for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=1)}

    rows = []
    for excel_row in ws.iter_rows(min_row=2):
        row = {"source": source, "source_label": source_label}
        for col in COLUMNS:
            idx = header.get(col)
            if not idx:
                row[col] = ""
                continue
            cell = excel_row[idx - 1]
            if col in HYPERLINK_COLUMNS and cell.hyperlink is not None:
                row[col] = clean(cell.hyperlink.target)
            else:
                row[col] = clean(cell.value)
        rows.append(row)
    return rows


def main():
    readme_summaries = {}
    if README_CACHE.exists():
        readme_summaries = json.loads(README_CACHE.read_text(encoding="utf-8"))

    all_rows: dict[str, dict] = {}
    for src in SOURCES:
        xlsx_path = src["xlsx"]
        if not xlsx_path.exists():
            print(f"Skipping {src['source']} — {xlsx_path.name} not found yet.")
            continue
        rows = load_rows(xlsx_path, src["source"], src["source_label"])
        print(f"{src['source']}: {len(rows)} rows from {xlsx_path.name}")
        for row in rows:
            key = row.get("repo_full_name", "")
            # A repo matched by both sourcers keeps whichever copy was loaded last;
            # order in SOURCES above decides precedence.
            if key:
                all_rows[key] = row

    rows = list(all_rows.values())
    for row in rows:
        row["readme_summary"] = readme_summaries.get(row.get("repo_full_name", ""), "")

    OUT.write_text("window.__DATA__ = " + json.dumps(rows, ensure_ascii=False) + ";", encoding="utf-8")
    with_summary = sum(1 for r in rows if r["readme_summary"])
    print(f"Wrote {len(rows)} unique rows -> {OUT} ({with_summary} with a README summary)")


if __name__ == "__main__":
    main()
