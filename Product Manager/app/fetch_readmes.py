"""
Fetches each candidate repo's README from GitHub and caches a short plain-text
summary in readme_summaries.json, keyed by repo_full_name. Run manually after
each sourcer run (before export_data.py): ../.venv/bin/python fetch_readmes.py

Only fetches repos not already in the cache, so re-runs are cheap.
"""
import json
import os
import re
import time
from pathlib import Path

import requests
from dotenv import load_dotenv

ROOT = Path(__file__).parent.parent
load_dotenv(ROOT / ".env")

XLSX = ROOT / "github_fde_react_candidates.xlsx"
CACHE_FILE = Path(__file__).parent / "readme_summaries.json"

GITHUB_TOKEN = (os.getenv("GITHUB_TOKEN") or "").strip()
SESSION = requests.Session()
SESSION.headers.update({
    "Accept": "application/vnd.github.raw+json",
    "User-Agent": "fde-react-sourcer-readme/1.0",
})
if GITHUB_TOKEN:
    SESSION.headers.update({"Authorization": f"Bearer {GITHUB_TOKEN}"})

TIMEOUT = 15
MAX_SUMMARY_CHARS = 320

_BADGE_LINE = re.compile(r"^\s*\[?!\[.*$")
_HEADER_LINE = re.compile(r"^\s*#{1,6}\s+")
_HTML_TAG = re.compile(r"<[^>]+>")
_MD_IMAGE = re.compile(r"!\[[^\]]*\]\([^)]*\)")
_MD_REF_BADGE = re.compile(r"\[!\[[^\]]*\]\[[^\]]*\]\]\[[^\]]*\]")
_MD_LINK = re.compile(r"\[([^\]]*)\]\([^)]*\)")
_MD_EMPHASIS = re.compile(r"[*_`>]+")


def load_repo_names() -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(XLSX, read_only=True)
    ws = wb.active
    header = {c.value: i for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=1)}
    idx = header.get("repo_full_name")
    names = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        val = row[idx - 1] if idx else None
        if val:
            names.append(str(val))
    return names


def clean_line(line: str) -> str:
    line = _MD_REF_BADGE.sub("", line)
    line = _MD_IMAGE.sub("", line)
    line = _MD_LINK.sub(r"\1", line)
    line = _HTML_TAG.sub("", line)
    line = _MD_EMPHASIS.sub("", line)
    return line.strip()


def summarize(readme_text: str) -> str:
    lines = readme_text.splitlines()
    paragraph_parts: list[str] = []
    for raw in lines:
        if _HEADER_LINE.match(raw) or _BADGE_LINE.match(raw):
            # Titles and badge rows aren't a description — skip them, don't let
            # them end the paragraph search either (a break here would mean
            # "no description" for every repo whose README opens with a badge row).
            continue
        cleaned = clean_line(raw)
        if not cleaned:
            if paragraph_parts:
                break
            continue
        if len(cleaned) < 15:
            continue
        paragraph_parts.append(cleaned)
        if sum(len(p) for p in paragraph_parts) > MAX_SUMMARY_CHARS:
            break

    text = " ".join(paragraph_parts).strip()
    if not text:
        return ""
    if len(text) <= MAX_SUMMARY_CHARS:
        return text

    truncated = text[:MAX_SUMMARY_CHARS]
    last_period = truncated.rfind(". ")
    if last_period > 60:
        return truncated[:last_period + 1]
    return truncated.rstrip() + "…"


def fetch_readme(full_name: str) -> str:
    url = f"https://api.github.com/repos/{full_name}/readme"
    for attempt in range(3):
        try:
            r = SESSION.get(url, timeout=TIMEOUT)
        except requests.RequestException:
            time.sleep(1.5 * (attempt + 1))
            continue
        if r.status_code == 200:
            return r.text
        if r.status_code == 404:
            return ""
        if r.status_code in (403, 429):
            time.sleep(3 * (attempt + 1))
            continue
        return ""
    return ""


def main():
    cache = {}
    if CACHE_FILE.exists():
        cache = json.loads(CACHE_FILE.read_text(encoding="utf-8"))

    names = load_repo_names()
    todo = [n for n in names if n not in cache]
    print(f"{len(names)} repos total, {len(todo)} need fetching (rest cached).")

    for i, full_name in enumerate(todo, 1):
        raw = fetch_readme(full_name)
        cache[full_name] = summarize(raw) if raw else ""
        if i % 20 == 0:
            CACHE_FILE.write_text(json.dumps(cache, ensure_ascii=False, indent=0), encoding="utf-8")
            print(f"  {i}/{len(todo)}")

    CACHE_FILE.write_text(json.dumps(cache, ensure_ascii=False, indent=0), encoding="utf-8")
    print(f"Done. Cache has {len(cache)} entries -> {CACHE_FILE}")


if __name__ == "__main__":
    main()
